import pandas as pd
from typing import Dict, List
import os
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import warnings
warnings.filterwarnings('ignore')


class Exporter:
    def __init__(self, output_dir: str = "output", filename_prefix: str = "房源数据"):
        self.output_dir = output_dir
        self.filename_prefix = filename_prefix
        os.makedirs(output_dir, exist_ok=True)

    def export_all(self, df: pd.DataFrame, stats: Dict = None) -> Dict[str, str]:
        exported_files = {}
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        if not df.empty:
            excel_file = self.export_excel(df, stats, timestamp)
            csv_file = self.export_csv(df, timestamp)
            exported_files['excel'] = excel_file
            exported_files['csv'] = csv_file

        if stats:
            pdf_file = self.export_pdf(df, stats, timestamp)
            html_file = self.export_html(df, stats, timestamp)
            exported_files['pdf'] = pdf_file
            exported_files['html'] = html_file

        return exported_files

    def export_excel(self, df: pd.DataFrame, stats: Dict = None, timestamp: str = None) -> str:
        if timestamp is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        filename = os.path.join(self.output_dir, f"{self.filename_prefix}_{timestamp}.xlsx")

        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            if not df.empty:
                df.to_excel(writer, sheet_name='房源列表', index=False)

                worksheet = writer.sheets['房源列表']
                self._adjust_column_width(worksheet, df)

            if stats:
                stats_df = self._create_stats_dataframe(stats)
                stats_df.to_excel(writer, sheet_name='统计信息', index=False)

                if '统计信息' in writer.sheets:
                    stats_worksheet = writer.sheets['统计信息']
                    self._adjust_column_width(stats_worksheet, stats_df)

        print(f"[导出] Excel文件已保存: {filename}")
        return filename

    def export_csv(self, df: pd.DataFrame, timestamp: str = None) -> str:
        if timestamp is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        filename = os.path.join(self.output_dir, f"{self.filename_prefix}_{timestamp}.csv")

        if not df.empty:
            df.to_csv(filename, index=False, encoding='utf-8-sig')
            print(f"[导出] CSV文件已保存: {filename}")
        else:
            print("[导出] 数据为空，跳过CSV导出")

        return filename

    def export_pdf(self, df: pd.DataFrame, stats: Dict, timestamp: str = None) -> str:
        if timestamp is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        filename = os.path.join(self.output_dir, f"{self.filename_prefix}_{timestamp}.pdf")

        doc = SimpleDocTemplate(
            filename,
            pagesize=landscape(A4),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )

        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            spaceAfter=30,
            alignment=1
        )

        elements.append(Paragraph("三亚房源数据报告", title_style))
        elements.append(Spacer(1, 12))

        if stats:
            elements.append(Paragraph("统计信息", styles['Heading2']))
            elements.append(Spacer(1, 12))

            stats_data = self._prepare_stats_for_pdf(stats)
            if stats_data:
                stats_table = Table(stats_data, colWidths=[200, 200])
                stats_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 10),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ]))
                elements.append(stats_table)
                elements.append(PageBreak())

        if not df.empty:
            elements.append(Paragraph("房源列表", styles['Heading2']))
            elements.append(Spacer(1, 12))

            display_cols = ['标题', '区域', '小区', '户型', '面积', '总价(万)', '单价', '来源']
            available_cols = [col for col in display_cols if col in df.columns]

            if available_cols:
                pdf_df = df[available_cols].head(50)

                table_data = [available_cols]
                for _, row in pdf_df.iterrows():
                    row_data = []
                    for col in available_cols:
                        value = row[col]
                        if pd.isna(value):
                            value = ''
                        else:
                            value = str(value)[:30]
                        row_data.append(value)
                    table_data.append(row_data)

                col_width = (landscape(A4)[0] - 60) / len(available_cols)
                col_widths = [col_width] * len(available_cols)

                table = Table(table_data, colWidths=col_widths)
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 8),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ]))
                elements.append(table)

                if len(df) > 50:
                    elements.append(Spacer(1, 12))
                    elements.append(Paragraph(f"注：仅显示前50条数据，共{len(df)}条数据", styles['Normal']))

        doc.build(elements)
        print(f"[导出] PDF文件已保存: {filename}")
        return filename

    def export_html(self, df: pd.DataFrame, stats: Dict, timestamp: str = None) -> str:
        if timestamp is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

        filename = os.path.join(self.output_dir, f"{self.filename_prefix}_{timestamp}.html")

        html_content = self._generate_html(df, stats)

        with open(filename, 'w', encoding='utf-8') as f:
            f.write(html_content)

        print(f"[导出] HTML文件已保存: {filename}")
        return filename

    def _generate_html(self, df: pd.DataFrame, stats: Dict) -> str:
        html = """
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>三亚房源数据报告</title>
    <style>
        body {
            font-family: 'Microsoft YaHei', Arial, sans-serif;
            margin: 0;
            padding: 20px;
            background-color: #f5f5f5;
        }
        .container {
            max-width: 1400px;
            margin: 0 auto;
            background-color: white;
            padding: 30px;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }
        h1 {
            text-align: center;
            color: #2c3e50;
            border-bottom: 3px solid #3498db;
            padding-bottom: 20px;
        }
        h2 {
            color: #34495e;
            border-left: 4px solid #3498db;
            padding-left: 15px;
            margin-top: 30px;
        }
        .stats-grid {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(250px, 1fr));
            gap: 20px;
            margin: 20px 0;
        }
        .stat-card {
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 20px;
            border-radius: 10px;
            text-align: center;
        }
        .stat-card h3 {
            margin: 0;
            font-size: 14px;
            opacity: 0.9;
        }
        .stat-card p {
            margin: 10px 0 0;
            font-size: 28px;
            font-weight: bold;
        }
        table {
            width: 100%;
            border-collapse: collapse;
            margin: 20px 0;
            font-size: 14px;
        }
        th {
            background-color: #3498db;
            color: white;
            padding: 12px;
            text-align: left;
        }
        td {
            padding: 10px;
            border-bottom: 1px solid #ddd;
        }
        tr:hover {
            background-color: #f5f5f5;
        }
        .price-low {
            color: #27ae60;
            font-weight: bold;
        }
        .price-medium {
            color: #f39c12;
            font-weight: bold;
        }
        .price-high {
            color: #e74c3c;
            font-weight: bold;
        }
        .tag {
            display: inline-block;
            background-color: #3498db;
            color: white;
            padding: 2px 8px;
            border-radius: 3px;
            font-size: 12px;
            margin-right: 5px;
        }
        .footer {
            text-align: center;
            margin-top: 30px;
            padding-top: 20px;
            border-top: 1px solid #ddd;
            color: #7f8c8d;
        }
        @media print {
            body {
                background-color: white;
            }
            .container {
                box-shadow: none;
            }
        }
    </style>
</head>
<body>
    <div class="container">
        <h1>🏠 三亚房源数据报告</h1>
"""
        if stats:
            html += """
        <h2>📊 数据统计</h2>
        <div class="stats-grid">
"""
            if '总房源数' in stats:
                html += f"""
            <div class="stat-card">
                <h3>总房源数</h3>
                <p>{stats['总房源数']}</p>
            </div>
"""
            if '平均总价(万)' in stats:
                html += f"""
            <div class="stat-card">
                <h3>平均总价</h3>
                <p>{stats['平均总价(万)']}万</p>
            </div>
"""
            if '平均面积(㎡)' in stats:
                html += f"""
            <div class="stat-card">
                <h3>平均面积</h3>
                <p>{stats['平均面积(㎡)']}㎡</p>
            </div>
"""
            if '爬取时间' in stats:
                html += f"""
            <div class="stat-card">
                <h3>爬取时间</h3>
                <p style="font-size: 16px;">{stats['爬取时间']}</p>
            </div>
"""
            html += "        </div>"

        if not df.empty:
            html += """
        <h2>📋 房源列表</h2>
        <table>
            <thead>
                <tr>
"""
            display_cols = ['标题', '区域', '小区', '户型', '面积', '总价(万)', '单价', '楼层', '来源']
            available_cols = [col for col in display_cols if col in df.columns]

            for col in available_cols:
                html += f"                    <th>{col}</th>\n"

            html += """                </tr>
            </thead>
            <tbody>
"""

            for _, row in df.iterrows():
                html += "                <tr>\n"
                for col in available_cols:
                    value = row[col]
                    if pd.isna(value):
                        value = ''
                    else:
                        value = str(value)

                    if col == '总价(万)':
                        try:
                            price = float(value)
                            if price < 100:
                                html += f'                    <td class="price-low">{value}</td>\n'
                            elif price < 150:
                                html += f'                    <td class="price-medium">{value}</td>\n'
                            else:
                                html += f'                    <td class="price-high">{value}</td>\n'
                        except:
                            html += f'                    <td>{value}</td>\n'
                    elif col == '标题':
                        html += f'                    <td><a href="{row.get("链接", "#")}" target="_blank">{value[:40]}</a></td>\n'
                    else:
                        html += f'                    <td>{value}</td>\n'

                html += "                </tr>\n"

            html += """            </tbody>
        </table>
"""

        html += f"""
        <div class="footer">
            <p>数据来源：链家网、安居客、58同城、贝壳找房 | 自动生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        </div>
    </div>
</body>
</html>
"""
        return html

    def _adjust_column_width(self, worksheet, df: pd.DataFrame):
        from openpyxl.utils import get_column_letter
        
        for idx, col in enumerate(df.columns, start=1):
            max_length = len(str(col))
            col_letter = get_column_letter(idx)
            
            for row in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=idx)
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[col_letter].width = adjusted_width

    def _create_stats_dataframe(self, stats: Dict) -> pd.DataFrame:
        rows = []
        for key, value in stats.items():
            if isinstance(value, dict):
                for sub_key, sub_value in value.items():
                    rows.append({'指标': f"{key} - {sub_key}", '数值': sub_value})
            else:
                rows.append({'指标': key, '数值': value})
        return pd.DataFrame(rows)

    def _prepare_stats_for_pdf(self, stats: Dict) -> List:
        data = []
        for key, value in stats.items():
            if not isinstance(value, dict):
                data.append([str(key), str(value)])
        return data
