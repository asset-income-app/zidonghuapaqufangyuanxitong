from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
from typing import Dict
from datetime import datetime
import os


class WordExporter:
    def __init__(self, output_dir: str = "output"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def export_to_word(self, df: pd.DataFrame, stats: Dict, 
                       filename: str = None) -> str:
        if filename is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = os.path.join(self.output_dir, f"房源报告_{timestamp}.docx")

        doc = Document()

        title = doc.add_heading('三亚房源数据分析报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph()

        doc.add_heading('一、数据概览', 1)

        table = doc.add_table(rows=6, cols=2)
        table.style = 'Light Grid Accent 1'

        cells = table.rows[0].cells
        cells[0].text = '总房源数'
        cells[1].text = str(stats.get('总房源数', 0))

        cells = table.rows[1].cells
        cells[0].text = '平均价格'
        cells[1].text = f"{stats.get('平均总价(万)', 0):.2f} 万元"

        cells = table.rows[2].cells
        cells[0].text = '最低价格'
        cells[1].text = f"{stats.get('最低总价(万)', 0):.2f} 万元"

        cells = table.rows[3].cells
        cells[0].text = '最高价格'
        cells[1].text = f"{stats.get('最高总价(万)', 0):.2f} 万元"

        cells = table.rows[4].cells
        cells[0].text = '平均面积'
        cells[1].text = f"{stats.get('平均面积(㎡)', 0):.2f} ㎡"

        cells = table.rows[5].cells
        cells[0].text = '数据来源'
        cells[1].text = '链家、安居客、58同城、贝壳'

        doc.add_paragraph()

        if '各区域房源数' in stats:
            doc.add_heading('二、区域分布', 1)
            for district, count in stats['各区域房源数'].items():
                doc.add_paragraph(f"{district}: {count} 套", style='List Bullet')
            doc.add_paragraph()

        if '各网站房源数' in stats:
            doc.add_heading('三、数据来源', 1)
            for source, count in stats['各网站房源数'].items():
                doc.add_paragraph(f"{source}: {count} 套", style='List Bullet')
            doc.add_paragraph()

        if not df.empty:
            doc.add_heading('四、房源列表（前50条）', 1)

            display_cols = ['标题', '区域', '户型', '面积', '总价(万)', '来源']
            available_cols = [col for col in display_cols if col in df.columns]

            if available_cols:
                table = doc.add_table(rows=min(51, len(df)+1), cols=len(available_cols))
                table.style = 'Light Grid Accent 1'

                header_cells = table.rows[0].cells
                for i, col in enumerate(available_cols):
                    header_cells[i].text = col
                    header_cells[i].paragraphs[0].runs[0].bold = True

                for idx, row in df.head(50).iterrows():
                    row_cells = table.rows[idx+1].cells
                    for i, col in enumerate(available_cols):
                        value = row[col]
                        if pd.isna(value):
                            value = ''
                        else:
                            value = str(value)[:30]
                        row_cells[i].text = value

        doc.add_paragraph()
        doc.add_paragraph('---')
        footer = doc.add_paragraph('本报告由三亚房源智能爬取系统自动生成')
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.save(filename)
        print(f"[Word导出] 文件已保存: {filename}")
        return filename
